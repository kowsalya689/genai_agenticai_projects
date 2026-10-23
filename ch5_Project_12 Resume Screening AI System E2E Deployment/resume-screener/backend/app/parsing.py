"""Task 1: read PDF / DOCX / TXT into plain text."""
import io
import os

try:
    import pymupdf as fitz  # PyMuPDF >= 1.24.3
except ImportError:  # older releases
    import fitz
from docx import Document

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def _pdf(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                parts.append(cell.text)
    return "\n".join(parts)


def _txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def extract_text(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    else:
        text = _txt(data)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if len(text.split()) < 15:
        raise ValueError("No readable text found (scanned PDF or empty file).")
    return text
